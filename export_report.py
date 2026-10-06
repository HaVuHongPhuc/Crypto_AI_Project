import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_full_quant_report():
    doc = Document()

    # Cấu hình lề trang (Standard 1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Tiêu đề báo cáo
    title = doc.add_heading('BÁO CÁO KIẾN TRÚC HỆ THỐNG 6-AGENT CRYPTO QUANT TRADING BOT (BTC/USDT)', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Tác giả: HaVuHongPhuc | Repository: Crypto_AI_Project\nVai trò: Senior Quantitative Trading Engineer & Python System Architect')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('—' * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Cấu trúc nội dung 10 Chương
    chapters = [
        ("CHƯƠNG I: TỔNG QUAN HỆ THỐNG VÀ KIẾN TRÚC TỔNG THỂ", 
         "Hệ thống vận hành theo mô hình Hybrid kết hợp giữa Machine Learning định lượng (Random Forest) và Hệ thống Đa tác tử (6-Agent LLM). Phân bổ vốn linh hoạt: 70 USDT cho chế độ có xu hướng mạnh (Trend Mode) và 30 USDT cho chế độ đi ngang (Flexible Mode) trên tổng vốn gốc 100 USDT."),
        
        ("CHƯƠNG II: BẢN ĐỒ CẤU TRÚC REPOSITORY VÀ QUẢN TRỊ FILE", 
         "Cấu trúc thư mục được module hóa cao độ: main.py điều hành nến 5m/1H; train.py huấn luyện mô hình cơ sở từ 35.000 nến; thư mục agents/ quản lý 6 tác tử; engine/ chịu trách nhiệm giả lập khớp lệnh và tái huấn luyện Walk-Forward."),
        
        ("CHƯƠNG III: MÔ HÌNH HỌC MÁY VÀ CHIẾN LƯỢC DỰ BÁO", 
         "Mô hình Random Forest (250 cây, max_depth=8, min_samples_leaf=15, class_weight='balanced') sử dụng 7 chỉ báo: RSI(14), MACD, Signal, EMA9, EMA21, EMA Spread %, Volume Change. Cửa sổ nhìn trước 4 nến (20 phút) với ngưỡng tăng trưởng >= 0.35% để gán nhãn BUY. Khắc phục triệt để lỗi đảo ngược nhãn bằng classes.index('BUY'). Cơ chế Walk-Forward tự động kích hoạt sau mỗi 2.016 nến (7 ngày) và Hot-Swap trực tiếp trong RAM."),
        
        ("CHƯƠNG IV: HẠ TẦNG LLM VÀ QUY TRÌNH ĐIỀU PHỐI ĐA TẦNG (WATERFALL)", 
         "Chuỗi dự phòng 3 tầng: Tier 1 (Gemini 1.5 Flash) -> Tier 2 (Ollama Cloud gemma4:31b) -> Tier 3 (Groq Cloud Qwen). Xử lý timeout/lỗi HTTP để chuyển tầng tự động. Quản lý token phân cấp: 500-600 token cho tác vụ thường nhật, 2000 token cho tác vụ tiến hóa quy tắc (auto_evolve_rules)."),
        
        ("CHƯƠNG V: TÌNH BÁO THỊ TRƯỜNG & PHÂN TÍCH TÂM LÝ (SENTIMENT AGENT)", 
         "Tự động cào dữ liệu RSS và Fear & Greed Index, trích xuất điểm hoảng loạn Panic Score từ 1 đến 10. Đưa ra 3 trạng thái chỉ thị: NORMAL, CAUTION, hoặc kích hoạt phanh khẩn cấp HALT_TRADING khi Panic Score >= 7 hoặc xuất hiện sự kiện Thiên nga đen."),
        
        ("CHƯƠNG VI: HỆ THỐNG ĐA TÁC TỬ VÀ QUY LUẬT ĐỒNG THUẬN", 
         "Quy trình đồng thuận 4 bước nghiêm ngặt: StrategistAgent khóa xu hướng 1H (ONLY_LONG / ONLY_SHORT) -> OperatorAgent đề xuất lệnh 5m dựa trên ML xác suất >= 65% hoặc < 35% -> SupervisorAgent thẩm định rủi ro và phủ quyết -> Thực thi."),
        
        ("CHƯƠNG VII: QUẢN TRỊ RỦI RO CƠ HỌC VÀ CƠ CHẾ THỰC THI", 
         "PaperTrader tự động trừ phí sàn 0.05%/chiều. Triển khai 3 tầng bảo vệ cơ học độc lập: Hard Stop-loss (-1.2%), Take Profit (+0.8% đến +1.5%), Trailing Stop (kích hoạt tại +0.4% với khoảng lùi 0.5%). Tự động đồng bộ và phục hồi trạng thái từ file active_position.json khi hệ thống khởi động lại."),
        
        ("CHƯƠNG VIII: PHẢN TƯ VÀ TIẾN HÓA ĐỘNG (REFLECTIVE LEARNING)", 
         "ReflectorAgent ghi nhận nhật ký vào memory.json sau mỗi lệnh đóng. Tự động kích hoạt hàm auto_evolve_rules khi gặp lệnh cắt lỗ (PnL < 0) hoặc định kỳ sau mỗi 5 lệnh để nâng cấp version bộ luật trong strategy_rules.json (v1 -> vN)."),
        
        ("CHƯƠNG IX: SRE, LOGGING VÀ GIÁM SÁT HỆ THỐNG (AUDITOR AGENT)", 
         "Ghi vết toàn diện vào bot.log. AuditorAgent đóng vai trò SRE tự động bắt các Exception, phân tích Traceback log, chẩn đoán sự cố mạng/API và đề xuất phương án tự chữa lành."),
        
        ("CHƯƠNG X: MA TRẬN SIÊU THAM SỐ VÀ LỘ TRÌNH PHÁT TRIỂN", 
         "Hệ thống sẵn sàng mở rộng sang giao dịch đa tài sản (ETH, SOL), tích hợp API Binance Live Trading và xây dựng mạng lưới bot học tập cộng tác.")
    ]

    for heading, body in chapters:
        doc.add_heading(heading, level=1)
        p = doc.add_paragraph(body)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(8)

    # Thêm bảng Ma trận siêu tham số
    doc.add_heading('BẢNG MA TRẬN SIÊU THAM SỐ HỆ THỐNG', level=2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Thành phần'
    hdr_cells[1].text = 'Tham số cấu hình'
    hdr_cells[2].text = 'Giá trị mặc định'

    params = [
        ("ML Model Core", "RF Estimators / Max Depth", "250 / 8"),
        ("ML Model Core", "Prediction Window (Look-ahead)", "4 nến (20 phút)"),
        ("ML Model Core", "Walk-Forward Retrain Period", "2.016 nến (7 ngày)"),
        ("Risk Manager", "Hard Stop-Loss", "-1.2%"),
        ("Risk Manager", "Trailing Stop Trigger / Gap", "+0.4% / 0.5%"),
        ("Risk Manager", "Trading Fee", "0.05% per trade"),
        ("Agent LLM", "Max Tokens (Standard / Evolve)", "600 / 2000 tokens"),
        ("Sentiment", "Panic Score Threshold", ">= 7 (HALT_TRADING)")
    ]

    for comp, param, val in params:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = param
        row_cells[2].text = val

    output_path = os.path.abspath("BaoCao_KienTruc_Crypto_AI_Project.docx")
    doc.save(output_path)
    print(f"SUCCESS: File Word đã được tạo thành công tại:\n{output_path}")

if __name__ == "__main__":
    create_full_quant_report()